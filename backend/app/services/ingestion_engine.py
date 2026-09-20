import os
import csv
import json
import openpyxl
from typing import List, Dict, Any
from app.core.security import sanitize_document_text, sanitize_filename
from app.services.ocr_service import OCRService

class IngestionEngine:
    @staticmethod
    def parse_file(file_path: str, original_filename: str) -> List[Dict[str, Any]]:
        """
        Ingests a file and returns parsed documents with page & provenance metadata.
        Supports PDF, DOCX, CSV, XLSX, JSON, TXT, and PNG/JPG/WEBP OCR.
        """
        ext = os.path.splitext(original_filename)[1].lower()
        docs = []

        # 1. Image Files -> OCR Engine
        if ext in [".png", ".jpg", ".jpeg", ".webp"]:
            return OCRService.process_image_file(file_path, original_filename)

        elif ext in [".txt", ".log"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            clean_content = sanitize_document_text(content)
            docs.append({
                "page_number": 1,
                "section_name": "Full Text",
                "raw_content": clean_content
            })

        elif ext == ".csv":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                rows = [", ".join(row) for row in reader if row]
            clean_content = sanitize_document_text("\n".join(rows))
            docs.append({
                "page_number": 1,
                "section_name": "CSV Data",
                "raw_content": clean_content
            })

        elif ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
                sheet = wb[sheet_name]
                lines = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val) for val in row if val is not None]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
                clean_content = sanitize_document_text("\n".join(lines))
                docs.append({
                    "page_number": sheet_idx,
                    "section_name": f"Sheet: {sheet_name}",
                    "raw_content": clean_content
                })

        elif ext == ".json":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                try:
                    data = json.load(f)
                    content = json.dumps(data, indent=2)
                except Exception:
                    f.seek(0)
                    content = f.read()
            clean_content = sanitize_document_text(content)
            docs.append({
                "page_number": 1,
                "section_name": "JSON Structured Data",
                "raw_content": clean_content
            })

        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(file_path)
                full_text = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        full_text.append(para.text.strip())
                clean_content = sanitize_document_text("\n".join(full_text))
                docs.append({
                    "page_number": 1,
                    "section_name": "Document Content",
                    "raw_content": clean_content
                })
            except Exception:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                docs.append({
                    "page_number": 1,
                    "section_name": "Fallback Document Text",
                    "raw_content": sanitize_document_text(content)
                })

        elif ext == ".pdf":
            pdf_docs = OCRService.process_scanned_pdf(file_path, original_filename)
            if pdf_docs:
                return pdf_docs

            # Basic text fallback if pdf parsing fails
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            docs.append({
                "page_number": 1,
                "section_name": "PDF Content",
                "raw_content": sanitize_document_text(content)
            })

        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            docs.append({
                "page_number": 1,
                "section_name": "Content",
                "raw_content": sanitize_document_text(content)
            })

        return docs
